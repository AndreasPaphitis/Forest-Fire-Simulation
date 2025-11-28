"""Design System for Professional Presentation"""

from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR

class Colors:
    UVA_RED = RGBColor(192, 0, 0)
    DARK_NAVY = RGBColor(0, 40, 85)
    MEDIUM_GRAY = RGBColor(100, 100, 100)
    LIGHT_GRAY = RGBColor(240, 240, 240)
    WHITE = RGBColor(255, 255, 255)
    ACCENT_GOLD = RGBColor(255, 200, 0)
    SUCCESS_GREEN = RGBColor(46, 125, 50)
    WARNING_ORANGE = RGBColor(255, 152, 0)
    NAVY_LIGHT = RGBColor(20, 60, 105)

class Fonts:
    TITLE_LARGE = Pt(44)
    TITLE_MEDIUM = Pt(36)
    TITLE_SMALL = Pt(32)
    SUBTITLE = Pt(24)
    BODY_LARGE = Pt(20)
    BODY_MEDIUM = Pt(18)
    BODY_SMALL = Pt(16)
    CAPTION = Pt(14)

class Layout:
    SLIDE_WIDTH = Inches(10)
    SLIDE_HEIGHT = Inches(7.5)
    MARGIN_LEFT = Inches(0.5)
    MARGIN_RIGHT = Inches(0.5)
    MARGIN_TOP = Inches(0.8)
    MARGIN_BOTTOM = Inches(0.5)
    CONTENT_WIDTH = SLIDE_WIDTH - MARGIN_LEFT - MARGIN_RIGHT
    CONTENT_HEIGHT = SLIDE_HEIGHT - MARGIN_TOP - MARGIN_BOTTOM
    HEADER_HEIGHT = Inches(0.8)
    SPACING_SMALL = Inches(0.2)
    SPACING_MEDIUM = Inches(0.4)
    ICON_SIZE = Inches(0.3)

def add_colored_header(slide, title_text, bg_color=Colors.UVA_RED):
    header = slide.shapes.add_shape(1, Inches(0), Inches(0), Layout.SLIDE_WIDTH, Layout.HEADER_HEIGHT)
    header.fill.solid()
    header.fill.fore_color.rgb = bg_color
    header.line.fill.background()
    
    title_box = slide.shapes.add_textbox(Layout.MARGIN_LEFT, Inches(0.15), Layout.CONTENT_WIDTH, Inches(0.5))
    text_frame = title_box.text_frame
    text_frame.word_wrap = True
    p = text_frame.paragraphs[0]
    p.text = title_text
    p.font.size = Fonts.TITLE_SMALL
    p.font.bold = True
    p.font.color.rgb = Colors.WHITE
    p.alignment = PP_ALIGN.LEFT
    return header, title_box

def add_accent_bar(slide, left, top, width, height, color=Colors.UVA_RED):
    bar = slide.shapes.add_shape(1, left, top, width, height)
    bar.fill.solid()
    bar.fill.fore_color.rgb = color
    bar.line.fill.background()
    return bar

def add_text_with_icon(slide, left, top, width, text, icon_color=Colors.UVA_RED, font_size=Fonts.BODY_LARGE):
    icon = slide.shapes.add_shape(2, left, top + Inches(0.05), Layout.ICON_SIZE, Layout.ICON_SIZE)
    icon.fill.solid()
    icon.fill.fore_color.rgb = icon_color
    icon.line.fill.background()
    
    text_box = slide.shapes.add_textbox(left + Layout.ICON_SIZE + Inches(0.15), top, width - Layout.ICON_SIZE - Inches(0.15), Inches(0.4))
    text_frame = text_box.text_frame
    text_frame.word_wrap = True
    p = text_frame.paragraphs[0]
    p.text = text
    p.font.size = font_size
    p.font.color.rgb = Colors.MEDIUM_GRAY
    return icon, text_box

def format_title_text(text_frame, text, color=Colors.DARK_NAVY, size=Fonts.TITLE_SMALL, bold=True):
    p = text_frame.paragraphs[0]
    p.text = text
    p.font.size = size
    p.font.bold = bold
    p.font.color.rgb = color
    p.alignment = PP_ALIGN.LEFT

def format_body_text(text_frame, text, color=Colors.MEDIUM_GRAY, size=Fonts.BODY_LARGE):
    p = text_frame.paragraphs[0]
    p.text = text
    p.font.size = size
    p.font.color.rgb = color
    p.alignment = PP_ALIGN.LEFT
