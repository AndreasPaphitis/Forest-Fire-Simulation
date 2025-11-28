"""Slide Layout Functions for Innovative Presentation"""

from pptx.util import Inches, Pt
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.enum.shapes import MSO_SHAPE
from pptx.dml.color import RGBColor
from design_system_innovative import (
    Colors, Fonts, Layout, add_dynamic_header, add_rounded_card,
    add_circular_overlay, create_gradient_background
)

def create_hero_title_slide(slide, title_text, subtitle_lines):
    """Create a hero-style title slide with full visual impact."""
    # Gradient background
    create_gradient_background(slide, Colors.DARK_NAVY, Colors.NAVY_LIGHT)
    
    # Large circular accent
    add_circular_overlay(slide, Inches(-1), Inches(-1), Inches(5), Colors.UVA_RED, 0.15)
    add_circular_overlay(slide, Inches(7), Inches(5), Inches(4), Colors.ACCENT_GOLD, 0.1)
    
    # Title
    title_box = slide.shapes.add_textbox(Inches(1), Inches(2), Inches(8), Inches(2))
    text_frame = title_box.text_frame
    text_frame.word_wrap = True
    text_frame.vertical_anchor = MSO_ANCHOR.MIDDLE
    p = text_frame.paragraphs[0]
    p.text = title_text
    p.font.size = Pt(40)
    p.font.bold = True
    p.font.color.rgb = Colors.WHITE
    p.alignment = PP_ALIGN.CENTER
    
    # Subtitles
    y_position = Inches(4.5)
    for line in subtitle_lines:
        subtitle_box = slide.shapes.add_textbox(Inches(1), y_position, Inches(8), Inches(0.3))
        text_frame = subtitle_box.text_frame
        p = text_frame.paragraphs[0]
        p.text = line
        p.font.size = Pt(18) if "Andreas" not in line else Pt(16)
        p.font.color.rgb = Colors.LIGHT_GRAY
        p.alignment = PP_ALIGN.CENTER
        y_position += Inches(0.35)

def create_split_content_slide(slide, title_text, content_items, key_stats=None):
    """Create a slide with diagonal split layout."""
    add_dynamic_header(slide, title_text, Colors.DARK_NAVY, use_diagonal=True)
    
    # Content area
    content_left = Inches(0.7)
    content_width = Inches(5.5)
    y_pos = Inches(1.5)
    
    for text, icon_color in content_items:
        if not text.strip():
            y_pos += Inches(0.25)
            continue
        
        # Rounded bullet
        bullet = slide.shapes.add_shape(
            MSO_SHAPE.OVAL,
            content_left, y_pos + Inches(0.05),
            Inches(0.25), Inches(0.25)
        )
        bullet.fill.solid()
        bullet.fill.fore_color.rgb = icon_color
        bullet.line.fill.background()
        
        # Text
        text_box = slide.shapes.add_textbox(
            content_left + Inches(0.4), y_pos,
            content_width, Inches(0.4)
        )
        text_frame = text_box.text_frame
        text_frame.word_wrap = True
        p = text_frame.paragraphs[0]
        p.text = text
        p.font.size = Fonts.BODY_MEDIUM
        p.font.color.rgb = Colors.MEDIUM_GRAY
        y_pos += Inches(0.45)
    
    # Key stats cards
    if key_stats:
        stats_left = Inches(6.7)
        y_pos = Inches(1.8)
        for label, value in key_stats:
            card = add_rounded_card(slide, stats_left, y_pos, Inches(2.8), Inches(0.9), Colors.LIGHT_GRAY)
            
            # Value
            value_box = slide.shapes.add_textbox(
                stats_left + Inches(0.2), y_pos + Inches(0.15),
                Inches(2.4), Inches(0.3)
            )
            text_frame = value_box.text_frame
            p = text_frame.paragraphs[0]
            p.text = value
            p.font.size = Fonts.SUBTITLE
            p.font.bold = True
            p.font.color.rgb = Colors.UVA_RED
            
            # Label
            label_box = slide.shapes.add_textbox(
                stats_left + Inches(0.2), y_pos + Inches(0.5),
                Inches(2.4), Inches(0.25)
            )
            text_frame = label_box.text_frame
            p = text_frame.paragraphs[0]
            p.text = label
            p.font.size = Fonts.BODY_SMALL
            p.font.color.rgb = Colors.MEDIUM_GRAY
            
            y_pos += Inches(1.1)

def create_image_overlay_slide(slide, title_text, content_text, figure_path):
    """Create a slide with image and overlapping content."""
    add_dynamic_header(slide, title_text, Colors.UVA_RED, use_diagonal=False)
    
    # Image
    if figure_path and figure_path.exists():
        try:
            slide.shapes.add_picture(
                str(figure_path),
                Inches(0.5), Inches(1.4),
                width=Inches(4.8)
            )
        except:
            pass
    
    # Content card overlapping image
    card = add_rounded_card(
        slide,
        Inches(5), Inches(1.6),
        Inches(4.5), Inches(5.4),
        Colors.WHITE
    )
    
    # Content text
    content_box = slide.shapes.add_textbox(
        Inches(5.3), Inches(1.9),
        Inches(4), Inches(5)
    )
    text_frame = content_box.text_frame
    text_frame.word_wrap = True
    
    for line in content_text.split('\n'):
        if line.strip():
            p = text_frame.add_paragraph() if text_frame.paragraphs[0].text else text_frame.paragraphs[0]
            p.text = line
            p.font.size = Fonts.BODY_MEDIUM
            p.font.color.rgb = Colors.MEDIUM_GRAY
            p.space_before = Pt(6)

def create_data_showcase_slide(slide, title_text, content_text, figure_path, metrics=None):
    """Create a slide focused on data visualization with metrics."""
    add_dynamic_header(slide, title_text, Colors.UVA_RED, use_diagonal=True)
    
    # Metrics circles in header
    if metrics:
        metric_x = Inches(6.5)
        for label, value, color in metrics[:3]:
            circle = slide.shapes.add_shape(
                MSO_SHAPE.OVAL,
                metric_x, Inches(0.15),
                Inches(0.6), Inches(0.6)
            )
            circle.fill.solid()
            circle.fill.fore_color.rgb = color
            circle.line.fill.background()
            
            # Value text
            value_box = slide.shapes.add_textbox(
                metric_x, Inches(0.25),
                Inches(0.6), Inches(0.2)
            )
            text_frame = value_box.text_frame
            text_frame.vertical_anchor = MSO_ANCHOR.MIDDLE
            p = text_frame.paragraphs[0]
            p.text = value
            p.font.size = Pt(14)
            p.font.bold = True
            p.font.color.rgb = Colors.WHITE
            p.alignment = PP_ALIGN.CENTER
            
            # Label
            label_box = slide.shapes.add_textbox(
                metric_x, Inches(0.42),
                Inches(0.6), Inches(0.2)
            )
            text_frame = label_box.text_frame
            p = text_frame.paragraphs[0]
            p.text = label
            p.font.size = Pt(9)
            p.font.color.rgb = Colors.WHITE
            p.alignment = PP_ALIGN.CENTER
            
            metric_x += Inches(0.9)
    
    # Content card
    card = add_rounded_card(
        slide,
        Inches(0.6), Inches(1.5),
        Inches(4.3), Inches(5.6),
        Colors.LIGHT_GRAY
    )
    
    content_box = slide.shapes.add_textbox(
        Inches(0.8), Inches(1.7),
        Inches(3.9), Inches(5.2)
    )
    text_frame = content_box.text_frame
    text_frame.word_wrap = True
    
    for line in content_text.split('\n'):
        if line.strip():
            p = text_frame.add_paragraph() if text_frame.paragraphs[0].text else text_frame.paragraphs[0]
            p.text = line
            p.font.size = Fonts.BODY_MEDIUM
            p.font.color.rgb = Colors.MEDIUM_GRAY
            p.space_before = Pt(4)
    
    # Figure
    if figure_path and figure_path.exists():
        try:
            slide.shapes.add_picture(
                str(figure_path),
                Inches(5.3), Inches(1.5),
                width=Inches(4.2)
            )
        except:
            pass

def create_comparison_cards_slide(slide, title_text, successes, warnings):
    """Create a comparison slide with card-based layout."""
    add_dynamic_header(slide, title_text, Colors.DARK_NAVY, use_diagonal=False)
    
    # Success card
    success_card = add_rounded_card(
        slide,
        Inches(0.6), Inches(1.5),
        Inches(4.3), Inches(5.5),
        Colors.LIGHT_GRAY
    )
    
    success_title = slide.shapes.add_textbox(
        Inches(0.9), Inches(1.8),
        Inches(3.7), Inches(0.4)
    )
    text_frame = success_title.text_frame
    p = text_frame.paragraphs[0]
    p.text = "Achievements"
    p.font.size = Fonts.SUBTITLE
    p.font.bold = True
    p.font.color.rgb = Colors.SUCCESS_GREEN
    
    y_pos = Inches(2.4)
    for success in successes:
        bullet = slide.shapes.add_shape(
            MSO_SHAPE.ROUNDED_RECTANGLE,
            Inches(1), y_pos,
            Inches(0.3), Inches(0.3)
        )
        bullet.fill.solid()
        bullet.fill.fore_color.rgb = Colors.SUCCESS_GREEN
        bullet.line.fill.background()
        
        text_box = slide.shapes.add_textbox(
            Inches(1.5), y_pos,
            Inches(3.2), Inches(0.5)
        )
        text_frame = text_box.text_frame
        text_frame.word_wrap = True
        p = text_frame.paragraphs[0]
        p.text = success
        p.font.size = Fonts.BODY_MEDIUM
        p.font.color.rgb = Colors.MEDIUM_GRAY
        y_pos += Inches(0.6)
    
    # Warnings card
    warning_card = add_rounded_card(
        slide,
        Inches(5.2), Inches(1.5),
        Inches(4.3), Inches(5.5),
        Colors.LIGHT_GRAY
    )
    
    warning_title = slide.shapes.add_textbox(
        Inches(5.5), Inches(1.8),
        Inches(3.7), Inches(0.4)
    )
    text_frame = warning_title.text_frame
    p = text_frame.paragraphs[0]
    p.text = "Challenges"
    p.font.size = Fonts.SUBTITLE
    p.font.bold = True
    p.font.color.rgb = Colors.WARNING_ORANGE
    
    y_pos = Inches(2.4)
    for warning in warnings:
        bullet = slide.shapes.add_shape(
            MSO_SHAPE.ROUNDED_RECTANGLE,
            Inches(5.6), y_pos,
            Inches(0.3), Inches(0.3)
        )
        bullet.fill.solid()
        bullet.fill.fore_color.rgb = Colors.WARNING_ORANGE
        bullet.line.fill.background()
        
        text_box = slide.shapes.add_textbox(
            Inches(6.1), y_pos,
            Inches(3.2), Inches(0.5)
        )
        text_frame = text_box.text_frame
        text_frame.word_wrap = True
        p = text_frame.paragraphs[0]
        p.text = warning
        p.font.size = Fonts.BODY_MEDIUM
        p.font.color.rgb = Colors.MEDIUM_GRAY
        y_pos += Inches(0.6)

def create_summary_overlay_slide(slide, title_text, sections):
    """Create a summary slide with overlapping elements."""
    add_dynamic_header(slide, title_text, Colors.UVA_RED, use_diagonal=True)
    
    y_pos = Inches(1.5)
    for number, heading, points in sections:
        # Number circle
        circle = slide.shapes.add_shape(
            MSO_SHAPE.OVAL,
            Inches(0.8), y_pos,
            Inches(0.5), Inches(0.5)
        )
        circle.fill.solid()
        circle.fill.fore_color.rgb = Colors.UVA_RED
        circle.line.fill.background()
        
        num_box = slide.shapes.add_textbox(
            Inches(0.8), y_pos,
            Inches(0.5), Inches(0.5)
        )
        text_frame = num_box.text_frame
        text_frame.vertical_anchor = MSO_ANCHOR.MIDDLE
        p = text_frame.paragraphs[0]
        p.text = str(number)
        p.font.size = Fonts.SUBTITLE
        p.font.bold = True
        p.font.color.rgb = Colors.WHITE
        p.alignment = PP_ALIGN.CENTER
        
        # Heading
        heading_box = slide.shapes.add_textbox(
            Inches(1.5), y_pos + Inches(0.05),
            Inches(7.8), Inches(0.4)
        )
        text_frame = heading_box.text_frame
        p = text_frame.paragraphs[0]
        p.text = heading
        p.font.size = Fonts.BODY_LARGE
        p.font.bold = True
        p.font.color.rgb = Colors.DARK_NAVY
        
        y_pos += Inches(0.5)
        
        # Points
        for point in points:
            point_box = slide.shapes.add_textbox(
                Inches(1.7), y_pos,
                Inches(7.6), Inches(0.3)
            )
            text_frame = point_box.text_frame
            text_frame.word_wrap = True
            p = text_frame.paragraphs[0]
            p.text = point
            p.font.size = Fonts.BODY_MEDIUM
            p.font.color.rgb = Colors.MEDIUM_GRAY
            y_pos += Inches(0.35)
        
        y_pos += Inches(0.3)

def create_closing_gradient_slide(slide, title_text, content_lines, style='acknowledgment'):
    """Create a closing slide with gradient background."""
    if style == 'thankyou':
        create_gradient_background(slide, Colors.DARK_NAVY, Colors.NAVY_LIGHT)
        
        add_circular_overlay(slide, Inches(-1), Inches(2), Inches(4), Colors.UVA_RED, 0.2)
        add_circular_overlay(slide, Inches(7), Inches(-1), Inches(5), Colors.ACCENT_GOLD, 0.15)
        
        title_box = slide.shapes.add_textbox(
            Inches(1), Inches(2.5),
            Inches(8), Inches(1.5)
        )
        text_frame = title_box.text_frame
        text_frame.vertical_anchor = MSO_ANCHOR.MIDDLE
        p = text_frame.paragraphs[0]
        p.text = title_text
        p.font.size = Fonts.TITLE_LARGE
        p.font.bold = True
        p.font.color.rgb = Colors.WHITE
        p.alignment = PP_ALIGN.CENTER
        
        y_pos = Inches(4.5)
        for line in content_lines:
            if line:
                line_box = slide.shapes.add_textbox(
                    Inches(1), y_pos,
                    Inches(8), Inches(0.4)
                )
                text_frame = line_box.text_frame
                p = text_frame.paragraphs[0]
                p.text = line
                p.font.size = Fonts.BODY_LARGE
                p.font.color.rgb = Colors.LIGHT_GRAY
                p.alignment = PP_ALIGN.CENTER
                y_pos += Inches(0.5)
    else:
        add_dynamic_header(slide, title_text, Colors.DARK_NAVY, use_diagonal=False)
        
        y_pos = Inches(2)
        for line in content_lines:
            if line.strip():
                line_box = slide.shapes.add_textbox(
                    Inches(2), y_pos,
                    Inches(6), Inches(0.35)
                )
                text_frame = line_box.text_frame
                p = text_frame.paragraphs[0]
                p.text = line
                if ':' in line and not line.startswith('•'):
                    p.font.size = Fonts.BODY_LARGE
                    p.font.bold = True
                    p.font.color.rgb = Colors.DARK_NAVY
                else:
                    p.font.size = Fonts.BODY_MEDIUM
                    p.font.color.rgb = Colors.MEDIUM_GRAY
                p.alignment = PP_ALIGN.CENTER
                y_pos += Inches(0.4)
