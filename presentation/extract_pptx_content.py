"""
Extract content from PowerPoint presentation
"""
from pptx import Presentation
from pathlib import Path

pptx_file = Path("presentation/Thesis_Presentation_Final_CITED.pptx")

print(f"Reading: {pptx_file}\n")
print("="*80)

prs = Presentation(str(pptx_file))

for i, slide in enumerate(prs.slides, 1):
    print(f"\n{'='*80}")
    print(f"SLIDE {i}")
    print('='*80)
    
    # Get title
    if slide.shapes.title:
        print(f"TITLE: {slide.shapes.title.text}")
        print()
    
    # Get all text from shapes
    texts = []
    for shape in slide.shapes:
        if hasattr(shape, "text") and shape.text.strip():
            if shape != slide.shapes.title:  # Skip title (already printed)
                texts.append(shape.text.strip())
    
    if texts:
        print("CONTENT:")
        for text in texts:
            print(f"  • {text}")
    
    # Check for images
    images = []
    for shape in slide.shapes:
        if shape.shape_type == 13:  # Picture
            images.append("Image present")
    
    if images:
        print(f"\n[{len(images)} image(s) on this slide]")

print("\n" + "="*80)
print(f"Total slides: {len(prs.slides)}")
print("="*80)






