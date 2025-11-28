"""Slide Layout Functions"""

from pptx.util import Inches, Pt
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.dml.color import RGBColor
from presentation.design_system import Colors, Fonts, Layout, add_colored_header, add_accent_bar, add_text_with_icon

def create_title_slide(slide, title_text, subtitle_lines):
    band = slide.shapes.add_shape(1, Inches(0), Inches(1.5), Layout.SLIDE_WIDTH, Inches(2))
    band.fill.solid()
    band.fill.fore_color.rgb = Colors.UVA_RED
    band.line.fill.background()
    
    title_box = slide.shapes.add_textbox(Inches(0.8), Inches(1.7), Inches(8.4), Inches(1.6))
    text_frame = title_box.text_frame
    text_frame.word_wrap = True
    text_frame.vertical_anchor = MSO_ANCHOR.MIDDLE
    p = text_frame.paragraphs[0]
    p.text = title_text
    p.font.size = Pt(38)
    p.font.bold = True
    p.font.color.rgb = Colors.WHITE
    p.alignment = PP_ALIGN.CENTER
    
    y_position = Inches(4)
    for line in subtitle_lines:
        subtitle_box = slide.shapes.add_textbox(Inches(1), y_position, Inches(8), Inches(0.3))
        text_frame = subtitle_box.text_frame
        p = text_frame.paragraphs[0]
        p.text = line
        p.font.size = Pt(18) if "Andreas" not in line else Pt(16)
        p.font.color.rgb = Colors.DARK_NAVY
        p.alignment = PP_ALIGN.CENTER
        y_position += Inches(0.35)
    
    bottom_bar = slide.shapes.add_shape(1, Inches(0), Inches(7.2), Layout.SLIDE_WIDTH, Inches(0.3))
    bottom_bar.fill.solid()
    bottom_bar.fill.fore_color.rgb = Colors.DARK_NAVY
    bottom_bar.line.fill.background()

def create_two_column_content_slide(slide, title_text, content_items, key_stats=None):
    add_colored_header(slide, title_text, Colors.DARK_NAVY)
    add_accent_bar(slide, Inches(0), Inches(0.8), Inches(0.08), Inches(6.7))
    
    content_left = Inches(0.6)
    content_width = Inches(5.5)
    y_pos = Inches(1.3)
    
    for text, color in content_items:
        if not text.strip():
            y_pos += Inches(0.3)
            continue
        add_text_with_icon(slide, content_left, y_pos, content_width, text, color)
        y_pos += Inches(0.5)
    
    if key_stats:
        stats_left = Inches(6.5)
        stats_width = Inches(3)
        y_pos = Inches(1.5)
        for label, value in key_stats:
            box = slide.shapes.add_shape(1, stats_left, y_pos, stats_width, Inches(0.7))
            box.fill.solid()
            box.fill.fore_color.rgb = Colors.LIGHT_GRAY
            box.line.color.rgb = Colors.MEDIUM_GRAY
            box.line.width = Pt(1)
            
            text_box = slide.shapes.add_textbox(stats_left + Inches(0.1), y_pos + Inches(0.1), stats_width - Inches(0.2), Inches(0.5))
            text_frame = text_box.text_frame
            text_frame.word_wrap = True
            text_frame.vertical_anchor = MSO_ANCHOR.MIDDLE
            p = text_frame.paragraphs[0]
            p.text = f"{label}\n{value}"
            p.font.size = Fonts.BODY_MEDIUM
            p.font.color.rgb = Colors.DARK_NAVY
            p.alignment = PP_ALIGN.LEFT
            y_pos += Inches(1.0)

def create_image_focused_slide(slide, title_text, content_text, figure_path):
    add_colored_header(slide, title_text, Colors.UVA_RED)
    
    if figure_path and figure_path.exists():
        try:
            slide.shapes.add_picture(str(figure_path), Inches(0.5), Inches(1.3), width=Inches(4.5))
        except:
            pass
    
    content_box = slide.shapes.add_textbox(Inches(5.3), Inches(1.5), Inches(4.2), Inches(5.5))
    text_frame = content_box.text_frame
    text_frame.word_wrap = True
    
    for line in content_text.split('\n'):
        if line.strip():
            p = text_frame.add_paragraph() if text_frame.paragraphs[0].text else text_frame.paragraphs[0]
            p.text = line
            p.font.size = Fonts.BODY_MEDIUM
            p.font.color.rgb = Colors.MEDIUM_GRAY
            p.space_before = Pt(6)

def create_structured_info_slide(slide, title_text, sections, figure_path=None):
    add_colored_header(slide, title_text, Colors.DARK_NAVY)
    add_accent_bar(slide, Inches(9.85), Inches(0.8), Inches(0.15), Inches(6.7))
    
    content_left = Inches(0.6)
    content_width = Inches(5) if figure_path else Inches(8.8)
    y_pos = Inches(1.2)
    
    for section_title, content_lines in sections:
        section_box = slide.shapes.add_textbox(content_left, y_pos, content_width, Inches(0.35))
        text_frame = section_box.text_frame
        p = text_frame.paragraphs[0]
        p.text = section_title
        p.font.size = Fonts.SUBTITLE
        p.font.bold = True
        p.font.color.rgb = Colors.UVA_RED
        y_pos += Inches(0.4)
        
        for line in content_lines:
            line_box = slide.shapes.add_textbox(content_left + Inches(0.2), y_pos, content_width - Inches(0.2), Inches(0.3))
            text_frame = line_box.text_frame
            p = text_frame.paragraphs[0]
            p.text = line
            p.font.size = Fonts.BODY_MEDIUM
            p.font.color.rgb = Colors.MEDIUM_GRAY
            y_pos += Inches(0.35)
        
        y_pos += Inches(0.3)
    
    if figure_path and figure_path.exists():
        try:
            slide.shapes.add_picture(str(figure_path), Inches(5.8), Inches(1.5), width=Inches(3.8))
        except:
            pass

def create_data_visualization_slide(slide, title_text, content_text, figure_path, metrics=None):
    add_colored_header(slide, title_text, Colors.UVA_RED)
    
    if metrics:
        metric_x = Inches(6.5)
        for label, value, color in metrics[:3]:
            circle = slide.shapes.add_shape(2, metric_x, Inches(0.12), Inches(0.65), Inches(0.65))
            circle.fill.solid()
            circle.fill.fore_color.rgb = color
            circle.line.fill.background()
            
            value_box = slide.shapes.add_textbox(metric_x, Inches(0.25), Inches(0.65), Inches(0.2))
            text_frame = value_box.text_frame
            text_frame.vertical_anchor = MSO_ANCHOR.MIDDLE
            p = text_frame.paragraphs[0]
            p.text = value
            p.font.size = Pt(14)
            p.font.bold = True
            p.font.color.rgb = Colors.WHITE
            p.alignment = PP_ALIGN.CENTER
            
            label_box = slide.shapes.add_textbox(metric_x, Inches(0.45), Inches(0.65), Inches(0.2))
            text_frame = label_box.text_frame
            p = text_frame.paragraphs[0]
            p.text = label
            p.font.size = Pt(10)
            p.font.color.rgb = Colors.WHITE
            p.alignment = PP_ALIGN.CENTER
            metric_x += Inches(1.0)
    
    content_box = slide.shapes.add_textbox(Inches(0.6), Inches(1.3), Inches(4.2), Inches(5.8))
    text_frame = content_box.text_frame
    text_frame.word_wrap = True
    
    for line in content_text.split('\n'):
        if line.strip():
            p = text_frame.add_paragraph() if text_frame.paragraphs[0].text else text_frame.paragraphs[0]
            p.text = line
            p.font.size = Fonts.BODY_MEDIUM
            p.font.color.rgb = Colors.MEDIUM_GRAY
            p.space_before = Pt(4)
    
    if figure_path and figure_path.exists():
        try:
            slide.shapes.add_picture(str(figure_path), Inches(5.2), Inches(1.3), width=Inches(4.3))
        except:
            pass

def create_comparison_slide(slide, title_text, successes, warnings):
    add_colored_header(slide, title_text, Colors.DARK_NAVY)
    add_accent_bar(slide, Inches(4.95), Inches(1.0), Pt(3), Inches(6.2))
    
    success_title = slide.shapes.add_textbox(Inches(0.6), Inches(1.2), Inches(4), Inches(0.4))
    text_frame = success_title.text_frame
    p = text_frame.paragraphs[0]
    p.text = "✓ Achievements"
    p.font.size = Fonts.SUBTITLE
    p.font.bold = True
    p.font.color.rgb = Colors.SUCCESS_GREEN
    
    y_pos = Inches(1.8)
    for success in successes:
        icon = slide.shapes.add_shape(2, Inches(0.8), y_pos, Inches(0.35), Inches(0.35))
        icon.fill.solid()
        icon.fill.fore_color.rgb = Colors.SUCCESS_GREEN
        icon.line.fill.background()
        
        text_box = slide.shapes.add_textbox(Inches(1.3), y_pos, Inches(3.3), Inches(0.5))
        text_frame = text_box.text_frame
        text_frame.word_wrap = True
        p = text_frame.paragraphs[0]
        p.text = success
        p.font.size = Fonts.BODY_MEDIUM
        p.font.color.rgb = Colors.MEDIUM_GRAY
        y_pos += Inches(0.6)
    
    warning_title = slide.shapes.add_textbox(Inches(5.2), Inches(1.2), Inches(4), Inches(0.4))
    text_frame = warning_title.text_frame
    p = text_frame.paragraphs[0]
    p.text = "⚠ Challenges"
    p.font.size = Fonts.SUBTITLE
    p.font.bold = True
    p.font.color.rgb = Colors.WARNING_ORANGE
    
    y_pos = Inches(1.8)
    for warning in warnings:
        icon = slide.shapes.add_shape(2, Inches(5.4), y_pos, Inches(0.35), Inches(0.35))
        icon.fill.solid()
        icon.fill.fore_color.rgb = Colors.WARNING_ORANGE
        icon.line.fill.background()
        
        text_box = slide.shapes.add_textbox(Inches(5.9), y_pos, Inches(3.7), Inches(0.5))
        text_frame = text_box.text_frame
        text_frame.word_wrap = True
        p = text_frame.paragraphs[0]
        p.text = warning
        p.font.size = Fonts.BODY_MEDIUM
        p.font.color.rgb = Colors.MEDIUM_GRAY
        y_pos += Inches(0.6)

def create_summary_slide(slide, title_text, sections):
    add_colored_header(slide, title_text, Colors.UVA_RED)
    
    y_pos = Inches(1.3)
    for number, heading, points in sections:
        circle = slide.shapes.add_shape(2, Inches(0.7), y_pos, Inches(0.5), Inches(0.5))
        circle.fill.solid()
        circle.fill.fore_color.rgb = Colors.UVA_RED
        circle.line.fill.background()
        
        num_box = slide.shapes.add_textbox(Inches(0.7), y_pos, Inches(0.5), Inches(0.5))
        text_frame = num_box.text_frame
        text_frame.vertical_anchor = MSO_ANCHOR.MIDDLE
        p = text_frame.paragraphs[0]
        p.text = str(number)
        p.font.size = Fonts.SUBTITLE
        p.font.bold = True
        p.font.color.rgb = Colors.WHITE
        p.alignment = PP_ALIGN.CENTER
        
        heading_box = slide.shapes.add_textbox(Inches(1.4), y_pos, Inches(8), Inches(0.4))
        text_frame = heading_box.text_frame
        p = text_frame.paragraphs[0]
        p.text = heading
        p.font.size = Fonts.BODY_LARGE
        p.font.bold = True
        p.font.color.rgb = Colors.DARK_NAVY
        
        y_pos += Inches(0.45)
        for point in points:
            point_box = slide.shapes.add_textbox(Inches(1.6), y_pos, Inches(7.8), Inches(0.3))
            text_frame = point_box.text_frame
            text_frame.word_wrap = True
            p = text_frame.paragraphs[0]
            p.text = point
            p.font.size = Fonts.BODY_MEDIUM
            p.font.color.rgb = Colors.MEDIUM_GRAY
            y_pos += Inches(0.32)
        
        y_pos += Inches(0.4)

def create_closing_slide(slide, title_text, content_lines, style='acknowledgment'):
    if style == 'thankyou':
        title_box = slide.shapes.add_textbox(Inches(1), Inches(2.5), Inches(8), Inches(1))
        text_frame = title_box.text_frame
        text_frame.vertical_anchor = MSO_ANCHOR.MIDDLE
        p = text_frame.paragraphs[0]
        p.text = title_text
        p.font.size = Fonts.TITLE_LARGE
        p.font.bold = True
        p.font.color.rgb = Colors.DARK_NAVY
        p.alignment = PP_ALIGN.CENTER
        
        y_pos = Inches(4.5)
        for line in content_lines:
            if line:
                line_box = slide.shapes.add_textbox(Inches(1), y_pos, Inches(8), Inches(0.4))
                text_frame = line_box.text_frame
                p = text_frame.paragraphs[0]
                p.text = line
                p.font.size = Fonts.BODY_LARGE
                p.font.color.rgb = Colors.MEDIUM_GRAY
                p.alignment = PP_ALIGN.CENTER
                y_pos += Inches(0.5)
    else:
        add_colored_header(slide, title_text, Colors.DARK_NAVY)
        y_pos = Inches(1.8)
        for line in content_lines:
            if line.strip():
                line_box = slide.shapes.add_textbox(Inches(2), y_pos, Inches(6), Inches(0.35))
                text_frame = line_box.text_frame
                p = text_frame.paragraphs[0]
                p.text = line
                if ':' in line and not line.startswith('•'):
                    p.font.size = Fonts.BODY_LARGE
                    p.font.bold = True
                    p.font.color.rgb = Colors.DARK_NAVY
                else:
                    p.font.size = Fonts.BODY_MEDIUM
                    p.font.color.rgb = Colors.MEDIUM_GRAY
                p.alignment = PP_ALIGN.CENTER
                y_pos += Inches(0.4)








